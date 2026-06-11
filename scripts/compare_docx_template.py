#!/usr/bin/env python3
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DEFAULT_TEMPLATE_INDICES = {
    "cover_title": 8,
    "topic": 10,
    "student_id": 17,
    "name": 18,
    "major": 19,
    "teacher": 20,
    "college": 21,
    "article_title": 24,
    "author": 25,
    "affiliation": 26,
    "abstract": 28,
    "keywords": 29,
    "h1": 31,
    "body_after_h1": 32,
    "h2": 35,
    "body_after_h2": 36,
    "refs_heading": 51,
    "ref_item": 53,
}


def cm(value):
    return round(value.cm, 3) if value is not None and hasattr(value, "cm") else None


def pt(value):
    return round(value.pt, 2) if value is not None and hasattr(value, "pt") else value


def first_text_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def para_signature(paragraph, include_run=False):
    pf = paragraph.paragraph_format
    sig = {
        "style": paragraph.style.name,
        "align": str(paragraph.alignment),
        "left": cm(pf.left_indent),
        "first": cm(pf.first_line_indent),
        "line": pt(pf.line_spacing),
        "before": pt(pf.space_before),
        "after": pt(pf.space_after),
    }
    if include_run:
        run = first_text_run(paragraph)
        font = run.font if run else None
        east_asia = None
        if run is not None and run._element.rPr is not None and run._element.rPr.rFonts is not None:
            east_asia = run._element.rPr.rFonts.get(qn("w:eastAsia"))
        sig.update(
            {
                "font": font.name if font else None,
                "eastAsia": east_asia,
                "size": pt(font.size) if font else None,
                "bold": font.bold if font else None,
                "color": str(font.color.rgb) if font and font.color.rgb else None,
            }
        )
    return sig


def section_signature(document):
    result = []
    for section in document.sections:
        result.append(
            {
                "top": cm(section.top_margin),
                "bottom": cm(section.bottom_margin),
                "left": cm(section.left_margin),
                "right": cm(section.right_margin),
                "header": cm(section.header_distance),
                "footer": cm(section.footer_distance),
                "width": cm(section.page_width),
                "height": cm(section.page_height),
            }
        )
    return result


def load_mapping(path):
    if not path:
        return None
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return {name: tuple(value) for name, value in data.items()}


def parse_roles(raw):
    if not raw:
        return None
    roles = [role.strip() for role in raw.split(",") if role.strip()]
    if not roles:
        raise SystemExit("--roles was provided but no role names were found.")
    return roles


def clean_text(text):
    return re.sub(r"\s+", " ", text.replace("\t", " ")).strip()


def find_first(paragraphs, predicate, start=0):
    for index in range(start, len(paragraphs)):
        if predicate(paragraphs[index], clean_text(paragraphs[index].text)):
            return index
    return None


def find_reference_item(paragraphs):
    refs = [index for index, para in enumerate(paragraphs) if re.match(r"^\[\d+\]", para.text.strip())]
    if len(refs) >= 2:
        return refs[1]
    return refs[0] if refs else None


def detect_target_indices(document):
    paragraphs = document.paragraphs
    indices = {
        "cover_title": find_first(paragraphs, lambda p, t: "课程报告" in t),
        "topic": find_first(paragraphs, lambda p, t: t.startswith("题 目") or t.startswith("题目")),
        "student_id": find_first(paragraphs, lambda p, t: t.startswith("学") and "号" in t),
        "name": find_first(paragraphs, lambda p, t: t.startswith("姓") and "名" in t),
        "major": find_first(paragraphs, lambda p, t: t.startswith("专") and "业" in t),
        "teacher": find_first(paragraphs, lambda p, t: t.startswith("指") and "教" in t and "师" in t),
        "college": find_first(paragraphs, lambda p, t: t.startswith("院") and ("系" in t or "所" in t)),
        "abstract": find_first(paragraphs, lambda p, t: t.startswith("摘")),
        "keywords": find_first(paragraphs, lambda p, t: t.startswith("关键词")),
        "h1": find_first(paragraphs, lambda p, t: p.style.name == "Heading 1" and t != "参考文献"),
        "h2": find_first(paragraphs, lambda p, t: p.style.name == "Heading 2"),
        "refs_heading": find_first(paragraphs, lambda p, t: t == "参考文献"),
        "ref_item": find_reference_item(paragraphs),
    }

    if indices["college"] is not None:
        indices["article_title"] = find_first(
            paragraphs,
            lambda p, t: bool(t) and p.alignment is not None and str(p.alignment) == "CENTER (1)",
            start=indices["college"] + 1,
        )
    else:
        indices["article_title"] = None

    if indices["article_title"] is not None:
        indices["author"] = find_first(paragraphs, lambda p, t: bool(t), start=indices["article_title"] + 1)
    else:
        indices["author"] = None

    if indices["author"] is not None:
        indices["affiliation"] = find_first(paragraphs, lambda p, t: bool(t), start=indices["author"] + 1)
    else:
        indices["affiliation"] = None

    if indices["h1"] is not None:
        indices["body_after_h1"] = find_first(paragraphs, lambda p, t: bool(t) and p.style.name != "Heading 1", start=indices["h1"] + 1)
    else:
        indices["body_after_h1"] = None

    if indices["h2"] is not None:
        indices["body_after_h2"] = find_first(paragraphs, lambda p, t: bool(t) and p.style.name != "Heading 2", start=indices["h2"] + 1)
    else:
        indices["body_after_h2"] = None

    return indices


def build_role_pairs(template_doc, target_doc, mapping, roles=None):
    if mapping is None:
        target_indices = detect_target_indices(target_doc)
        role_pairs = {
            role: (template_index, target_indices.get(role))
            for role, template_index in DEFAULT_TEMPLATE_INDICES.items()
        }
    else:
        role_pairs = mapping

    if roles:
        missing = [role for role in roles if role not in role_pairs]
        if missing:
            available = ", ".join(role_pairs)
            raise SystemExit(f"Unknown role(s) in --roles: {', '.join(missing)}. Available: {available}")
        role_pairs = {role: role_pairs[role] for role in roles}

    return role_pairs


def role_map(role_pairs):
    return {
        role: {
            "template_index": pair[0],
            "target_index": pair[1],
        }
        for role, pair in role_pairs.items()
    }


def compare(template_doc, target_doc, role_pairs, include_run=False):
    failed = []
    for role, (template_index, target_index) in role_pairs.items():
        if template_index is None or target_index is None:
            failed.append(
                {
                    "role": role,
                    "error": "could not locate template or target paragraph",
                    "template_index": template_index,
                    "target_index": target_index,
                }
            )
            continue
        try:
            template_para = template_doc.paragraphs[template_index]
            target_para = target_doc.paragraphs[target_index]
        except IndexError as exc:
            failed.append({"role": role, "error": f"paragraph index out of range: {exc}"})
            continue
        expected = para_signature(template_para, include_run=include_run)
        actual = para_signature(target_para, include_run=include_run)
        diff = {key: [expected[key], actual[key]] for key in expected if expected[key] != actual[key]}
        if diff:
            failed.append(
                {
                    "role": role,
                    "template_index": template_index,
                    "target_index": target_index,
                    "template_text": template_para.text[:80],
                    "target_text": target_para.text[:80],
                    "diff": diff,
                }
            )
    return failed


def document_colors(docx_path):
    colors = set()
    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    colors.update(re.findall(r'w:color[^>]+', xml))
    return sorted(colors)


def build_report(args):
    template_path = Path(args.template)
    target_path = Path(args.target)
    template_doc = Document(str(template_path))
    target_doc = Document(str(target_path))
    mapping = load_mapping(args.mapping)
    roles = parse_roles(args.roles)
    pairs = build_role_pairs(template_doc, target_doc, mapping, roles=roles)

    expected_sections = section_signature(template_doc)
    actual_sections = section_signature(target_doc)
    section_failed = []
    if expected_sections != actual_sections:
        section_failed.append({"expected": expected_sections, "actual": actual_sections})

    format_failed = compare(template_doc, target_doc, pairs, include_run=args.include_run)

    colors = None
    color_failed = []
    if args.check_colors or args.require_black_only:
        colors = document_colors(target_path)
        if args.require_black_only:
            color_failed = [
                color
                for color in colors
                if 'w:val="000000"' not in color and 'w:val="auto"' not in color
            ]

    return {
        "ok": not (section_failed or format_failed or color_failed),
        "template": str(template_path),
        "target": str(target_path),
        "mapping_source": str(args.mapping) if args.mapping else "auto",
        "roles": list(pairs),
        "role_map": role_map(pairs),
        "section_failed": section_failed,
        "format_failed": format_failed,
        "document_colors": colors,
        "color_failed": color_failed,
    }


def print_report(report, args):
    print("section_failed", json.dumps(report["section_failed"], ensure_ascii=False))
    print("format_failed", json.dumps(report["format_failed"], ensure_ascii=False))

    if args.dump_map:
        print("role_map", json.dumps(report["role_map"], ensure_ascii=False))

    if args.check_colors or args.require_black_only:
        print("document_colors", json.dumps(report["document_colors"], ensure_ascii=False))
        if args.require_black_only:
            print("color_failed", json.dumps(report["color_failed"], ensure_ascii=False))

    if args.summary:
        print(
            "summary",
            json.dumps(
                {
                    "ok": report["ok"],
                    "section_failures": len(report["section_failed"]),
                    "format_failures": len(report["format_failed"]),
                    "color_failures": len(report["color_failed"]),
                },
                ensure_ascii=False,
            ),
        )


def write_report(report, path):
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="Compare a target DOCX against a Word template.")
    parser.add_argument("--template", required=True, help="Template .docx path")
    parser.add_argument("--target", required=True, help="Target .docx path")
    parser.add_argument("--mapping", help="JSON mapping of role to [template_index, target_index]")
    parser.add_argument("--roles", help="Comma-separated role names to compare, such as h1,h2,ref_item")
    parser.add_argument("--include-run", action="store_true", help="Also compare first text run font properties")
    parser.add_argument("--check-colors", action="store_true", help="Report document XML color values")
    parser.add_argument("--require-black-only", action="store_true", help="Fail if target document has non-black explicit colors")
    parser.add_argument("--dump-map", action="store_true", help="Print the resolved role-to-paragraph mapping")
    parser.add_argument("--summary", action="store_true", help="Print a compact pass/fail summary")
    parser.add_argument("--report", help="Write a JSON compliance report to this path")
    args = parser.parse_args()

    report = build_report(args)
    print_report(report, args)
    if args.report:
        write_report(report, args.report)

    if not report["ok"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
